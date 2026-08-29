import os
import firebase_admin
from firebase_admin import credentials, firestore
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

cred = credentials.Certificate('serviceAccountKey.json')
firebase_admin.initialize_app(cred)
db = firestore.client()

users = list(db.collection('users').stream())
students_26 = []

for u in users:
    d = u.to_dict()
    r = u.id.strip().upper()
    user_type = (d.get('userType') or 'student').lower().strip()
    
    if r.startswith('26') and user_type == 'student':
        wings = d.get('wings', [])
        if not isinstance(wings, list):
            wings = [wings] if wings else []
        
        wing_str = ", ".join(wings) if wings else "Unassigned"
        
        students_26.append({
            "roll": r,
            "name": d.get('name', '').strip(),
            "email": d.get('instituteOutlookId', '').strip(),
            "wing": wing_str,
            "wings_list": wings
        })

students_26.sort(key=lambda x: x['roll'])
print(f"Total 2026 batch students found: {len(students_26)}")

# Only individual wing configurations (No Summary Dashboard, No All Students)
WING_CONFIGS = [
    {
        "tab_name": "Teaching & Technical",
        "title": "Teaching and Technical Wing (TTW)",
        "color": "203764", # Deep Blue
        "filter_fn": lambda s: "Teaching and Technical Wing" in s['wings_list']
    },
    {
        "tab_name": "Environmental Wing",
        "title": "Environmental Wing (ENV)",
        "color": "375623", # Forest Green
        "filter_fn": lambda s: "Environmental Wing" in s['wings_list']
    },
    {
        "tab_name": "Prerna Wing",
        "title": "Prerna Wing (PRN)",
        "color": "C65911", # Rust Orange
        "filter_fn": lambda s: "Prerna Wing" in s['wings_list']
    },
    {
        "tab_name": "Rural Development",
        "title": "Rural Development Wing (RDW)",
        "color": "833C0C", # Warm Brown / Earth
        "filter_fn": lambda s: "Rural Development Wing" in s['wings_list']
    },
    {
        "tab_name": "Design & Curation",
        "title": "Design and Curation Wing (DCW)",
        "color": "7030A0", # Royal Purple
        "filter_fn": lambda s: "Design and Curation Wing" in s['wings_list']
    }
]

wb = openpyxl.Workbook()
wb.remove(wb.active) # Remove default sheet

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
headers = ["S.No", "Roll Number", "Student Name", "Institute Outlook ID", "Assigned Wing"]

for cfg in WING_CONFIGS:
    filtered_students = [s for s in students_26 if cfg["filter_fn"](s)]
    ws = wb.create_sheet(title=cfg["tab_name"])
    ws.views.sheetView[0].showGridLines = True
    
    # Page Title Banner
    ws.merge_cells("A1:E1")
    banner = ws["A1"]
    banner.value = f"{cfg['title']} ({len(filtered_students)} Students)"
    banner.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    banner.fill = PatternFill(start_color=cfg["color"], end_color=cfg["color"], fill_type="solid")
    banner.alignment = align_center
    ws.row_dimensions[1].height = 32

    # Column Headers (Row 3)
    ws.append([]) # Row 2 spacer
    ws.append(headers) # Row 3
    ws.row_dimensions[3].height = 24
    
    h_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    h_fill = PatternFill(start_color=cfg["color"], end_color=cfg["color"], fill_type="solid")
    
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=3, column=col_idx)
        cell.font = h_font
        cell.fill = h_fill
        cell.alignment = align_center
        cell.border = border_thin
    
    # Data Rows
    for i, s in enumerate(filtered_students, start=1):
        row_data = [i, s['roll'], s['name'], s['email'], s['wing']]
        ws.append(row_data)
        current_row = i + 3
        ws.row_dimensions[current_row].height = 20
        
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=current_row, column=col_idx)
            c.border = border_thin
            c.font = Font(name="Calibri", size=10)
            if col_idx in [1, 2]:
                c.alignment = align_center
            else:
                c.alignment = align_left
            
            if i % 2 == 0:
                c.fill = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
    
    # Enable AutoFilter on header row
    ws.auto_filter.ref = f"A3:E{len(filtered_students) + 3}"

# Auto-adjust column widths across all sheets
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        for cell in col:
            val_str = str(cell.value or '')
            if cell.row == 1: # Ignore title banner
                continue
            if len(val_str) > max_len:
                max_len = len(val_str)
        col_letter = get_column_letter(col[0].column)
        sheet.column_dimensions[col_letter].width = max(max_len + 5, 12)

excel_path = "scripts/nss_iitp_2026_wings_only.xlsx"
wb.save(excel_path)
print(f"Successfully saved wing-wise only Excel file to: {os.path.abspath(excel_path)}")
